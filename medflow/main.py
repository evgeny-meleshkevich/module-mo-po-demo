from __future__ import annotations

import io
import hmac
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader, select_autoescape
from PIL import Image
from openpyxl import load_workbook

from .auth import csrf_token, hash_password, new_token, token_digest, verify_password
from .catalog import clear_catalog_cache, get_catalog
from .config import APPROVED_CATALOG_HASH_PATH, BUSINESS_UNITS, CATALOG_STORAGE_DIR, ROOT_DIR, SIGNATURE_DIR
from .documents import generate_documents
from .models import BusinessUnitPreference, DocumentRequest, InvitationRequest, LoginRequest
from .storage import (
    count_active_admins,
    create_invitation,
    create_session,
    delete_account,
    delete_document,
    delete_invitation,
    delete_session,
    get_document,
    get_invitation,
    get_profile,
    get_profile_by_email,
    get_session_profile,
    get_catalog_version,
    initialize_storage,
    list_accounts,
    list_documents,
    list_invitations,
    list_catalog_versions,
    resolve_archived_file,
    save_account,
    save_catalog_version,
    save_document,
    save_last_business_unit,
    refresh_invitation,
    use_invitation,
    activate_catalog_version,
    upgrade_profile_to_account,
)
from .integrity import file_sha256


app = FastAPI(title="Направления МО и ПО", docs_url=None, redoc_url=None, openapi_url=None)
app.mount("/static", StaticFiles(directory=ROOT_DIR / "static"), name="static")
templates = Environment(
    loader=FileSystemLoader(ROOT_DIR / "templates"),
    autoescape=select_autoescape(["html", "xml"]),
)


@app.on_event("startup")
def startup() -> None:
    initialize_storage()


@app.get("/", response_class=HTMLResponse)
def index(request: Request) -> HTMLResponse:
    template = templates.get_template("index.html")
    return HTMLResponse(template.render())


SESSION_COOKIE = "medflow_test_session"
PUBLIC_API_ROUTES = {
    ("GET", "/api/auth/status"),
    ("POST", "/api/auth/login"),
    ("POST", "/api/auth/setup-admin"),
}


def _is_public_api_route(method: str, path: str) -> bool:
    if (method, path) in PUBLIC_API_ROUTES:
        return True
    # Обладателю действующей одноразовой ссылки нужен доступ к регистрации до входа.
    if method == "GET" and re.fullmatch(r"/api/invitations/[^/]+", path):
        return True
    return bool(method == "POST" and re.fullmatch(r"/api/invitations/[^/]+/accept", path))


@app.middleware("http")
async def require_session_for_api(request: Request, call_next):
    path = request.url.path
    if path.startswith("/api/"):
        if not _is_public_api_route(request.method, path):
            profile = _current_profile(request, required=False)
            if not profile:
                return JSONResponse(
                    status_code=401,
                    content={"detail": "Войдите в профиль"},
                    headers={"Cache-Control": "no-store"},
                )
            if request.method not in {"GET", "HEAD", "OPTIONS"}:
                expected = csrf_token(request.cookies.get(SESSION_COOKIE, ""))
                provided = request.headers.get("X-CSRF-Token", "")
                if not hmac.compare_digest(provided, expected):
                    return JSONResponse(
                        status_code=403,
                        content={"detail": "Проверка безопасности запроса не пройдена"},
                        headers={"Cache-Control": "no-store"},
                    )
    response = await call_next(request)
    if path.startswith("/api/") or path == "/":
        response.headers["Cache-Control"] = "no-store"
    return response


def _current_profile(request: Request, required: bool = True) -> dict | None:
    session_token = request.cookies.get(SESSION_COOKIE, "")
    if session_token:
        profile = get_session_profile(token_digest(session_token))
        if profile:
            return profile
    if required:
        raise HTTPException(401, "Войдите в профиль")
    return None


def _require_admin(request: Request) -> dict:
    profile = _current_profile(request)
    if profile["role"] != "admin":
        raise HTTPException(403, "Доступно только администратору")
    return profile


def _require_catalog_editor(request: Request) -> dict:
    profile = _current_profile(request)
    if profile["role"] not in {"admin", "editor"}:
        raise HTTPException(403, "Доступно только редактору или администратору")
    return profile


def _set_session(response: Response, profile_id: str) -> None:
    token = new_token()
    create_session(profile_id, token_digest(token))
    response.set_cookie(
        SESSION_COOKIE,
        token,
        max_age=30 * 24 * 60 * 60,
        httponly=True,
        samesite="lax",
        secure=False,
    )


def _profile_response(profile: dict) -> dict:
    return {
        "configured": True,
        "profile_id": profile["id"],
        "email": profile.get("email") or "",
        "display_name": profile.get("display_name") or profile.get("signer_name") or "",
        "role": profile["role"],
        "signer_name": profile["signer_name"],
        "has_signature": bool(profile.get("signature_path")),
        "last_business_unit_id": profile.get("last_business_unit_id", ""),
    }


@app.get("/api/profile")
def profile_get(request: Request) -> dict:
    profile = _current_profile(request, required=False)
    if not profile:
        return {"configured": False}
    return _profile_response(profile)


@app.get("/api/auth/status")
def auth_status(request: Request) -> dict:
    profile = _current_profile(request, required=False)
    return {
        "authenticated": bool(profile),
        "first_admin_required": count_active_admins() == 0,
        "profile": _profile_response(profile) if profile else None,
        "csrf_token": csrf_token(request.cookies[SESSION_COOKIE]) if profile else None,
    }


@app.post("/api/auth/setup-admin")
def setup_first_admin(payload: LoginRequest, request: Request, response: Response) -> dict:
    if count_active_admins() != 0:
        raise HTTPException(409, "Первый администратор уже создан")
    email = payload.email.strip().lower()
    if get_profile_by_email(email):
        raise HTTPException(409, "Аккаунт с таким email уже существует")
    current = _current_profile(request, required=False)
    if current and not current.get("password_hash"):
        profile = upgrade_profile_to_account(
            current["id"], email=email, role="admin", display_name=email,
            password_hash=hash_password(payload.password),
        )
    else:
        profile = save_account(
            email=email,
            role="admin",
            display_name=email,
            signer_name="",
            signature_path=None,
            password_hash=hash_password(payload.password),
        )
    _set_session(response, profile["id"])
    return _profile_response(profile)


@app.post("/api/auth/login")
def login(payload: LoginRequest, response: Response) -> dict:
    profile = get_profile_by_email(payload.email)
    if not profile or not profile.get("password_hash") or not verify_password(payload.password, profile["password_hash"]):
        raise HTTPException(401, "Неверный email или пароль")
    _set_session(response, profile["id"])
    return _profile_response(profile)


@app.post("/api/auth/logout")
def logout(request: Request, response: Response) -> dict:
    token = request.cookies.get(SESSION_COOKIE, "")
    if token:
        delete_session(token_digest(token))
    response.delete_cookie(SESSION_COOKIE)
    return {"ok": True}


@app.get("/api/invitations/{token}")
def invitation_get(token: str) -> dict:
    invitation = get_invitation(token_digest(token))
    if not invitation or invitation["used_at"] or invitation["expires_at"] <= datetime.now().isoformat():
        raise HTTPException(404, "Приглашение недействительно или срок его действия истёк")
    return {"email": invitation["email"], "role": invitation["role"]}


@app.post("/api/invitations/{token}/accept")
async def invitation_accept(
    token: str,
    response: Response,
    password: str = Form(...),
    display_name: str = Form(""),
    signer_name: str = Form(""),
    signature: UploadFile | None = File(None),
) -> dict:
    invitation = get_invitation(token_digest(token))
    if not invitation or invitation["used_at"] or invitation["expires_at"] <= datetime.now().isoformat():
        raise HTTPException(404, "Приглашение недействительно или срок его действия истёк")
    if len(password) < 8:
        raise HTTPException(400, "Пароль должен содержать не менее 8 символов")
    if get_profile_by_email(invitation["email"]):
        raise HTTPException(409, "Аккаунт с таким email уже существует")
    role = invitation["role"]
    signer_name = signer_name.strip()
    signature_path = None
    profile_id = uuid4().hex
    if role == "user":
        if not signer_name:
            raise HTTPException(400, "Укажите фамилию и инициалы")
        if not signature or not signature.filename:
            raise HTTPException(400, "Загрузите изображение подписи")
        contents = await signature.read()
        try:
            image = Image.open(io.BytesIO(contents)).convert("RGBA")
        except Exception as exc:
            raise HTTPException(400, "Файл подписи должен быть изображением") from exc
        target = SIGNATURE_DIR / f"{profile_id}.png"
        image.save(target, "PNG")
        signature_path = str(target)
    profile = save_account(
        email=invitation["email"],
        role=role,
        display_name=display_name.strip() or signer_name or invitation["email"],
        signer_name=signer_name if role == "user" else "",
        signature_path=signature_path,
        password_hash=hash_password(password),
        profile_id=profile_id,
    )
    use_invitation(invitation["id"])
    _set_session(response, profile["id"])
    return _profile_response(profile)


@app.get("/api/admin/users")
def admin_users(request: Request) -> list[dict]:
    _require_admin(request)
    return list_accounts()


@app.delete("/api/admin/users/{profile_id}")
def admin_user_delete(profile_id: str, request: Request) -> dict:
    admin = _require_admin(request)
    if profile_id == admin["id"]:
        raise HTTPException(400, "Нельзя удалить собственный аккаунт")
    deleted_documents = delete_account(profile_id)
    if deleted_documents is None:
        raise HTTPException(400, "Администраторов удалять нельзя")
    return {"ok": True, "deleted_documents": deleted_documents}


@app.get("/api/admin/invitations")
def admin_invitations(request: Request) -> list[dict]:
    _require_admin(request)
    return list_invitations()


@app.post("/api/admin/invitations")
def admin_invitation_create(payload: InvitationRequest, request: Request) -> dict:
    admin = _require_admin(request)
    email = payload.email.strip().lower()
    if get_profile_by_email(email):
        raise HTTPException(409, "Аккаунт с таким email уже существует")
    token = new_token()
    invitation = create_invitation(email, payload.role, token_digest(token), admin["id"])
    return {
        "id": invitation["id"],
        "email": email,
        "role": payload.role,
        "expires_at": invitation["expires_at"],
        "invite_url": f"{str(request.base_url).rstrip('/')}?invite={token}",
    }


@app.post("/api/admin/invitations/{invitation_id}/link")
def admin_invitation_link(invitation_id: str, request: Request) -> dict:
    _require_admin(request)
    token = new_token()
    invitation = refresh_invitation(invitation_id, token_digest(token))
    if not invitation:
        raise HTTPException(404, "Ожидающее приглашение не найдено")
    return {
        "invite_url": f"{str(request.base_url).rstrip('/')}?invite={token}",
        "expires_at": invitation["expires_at"],
    }


@app.delete("/api/admin/invitations/{invitation_id}")
def admin_invitation_delete(invitation_id: str, request: Request) -> dict:
    _require_admin(request)
    if not delete_invitation(invitation_id):
        raise HTTPException(404, "Приглашение не найдено")
    return {"ok": True}


def _validate_catalog_file(path: Path) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    counts: dict[str, int] = {}
    combinations: dict[str, set[tuple[str, str]]] = {}
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"valid": False, "errors": ["Не удалось открыть Excel-файл"], "warnings": [], "counts": {}}
    try:
        for form_name in ("Направления_на_МО", "Направления на ПО (автомат!!!)"):
            if form_name not in workbook.sheetnames:
                errors.append(f"В единой демонстрационной книге отсутствует форма «{form_name}»")
        for sheet_name in ("Конт МО", "Конт ПО"):
            if sheet_name not in workbook.sheetnames:
                errors.append(f"Отсутствует лист «{sheet_name}»")
                continue
            sheet = workbook[sheet_name]
            seen: set[tuple[str, str]] = set()
            filled = 0
            duplicates = 0
            for row in range(4, sheet.max_row + 1):
                unit = str(sheet.cell(row, 2).value or "").strip()
                job = str(sheet.cell(row, 3).value or "").strip()
                requirement = str(sheet.cell(row, 4).value or "").strip()
                if not job and not requirement:
                    continue
                if not unit or not job or not requirement:
                    errors.append(f"Лист «{sheet_name}», строка {row}: заполнены не все обязательные ячейки")
                    continue
                key = (unit.casefold(), job.casefold())
                if key in seen:
                    duplicates += 1
                seen.add(key)
                filled += 1
            counts[sheet_name] = filled
            combinations[sheet_name] = seen
            if filled == 0:
                errors.append(f"На листе «{sheet_name}» нет записей")
            if duplicates:
                warnings.append(f"Лист «{sheet_name}»: найдено повторяющихся комбинаций — {duplicates}")
        if all(name in combinations for name in ("Конт МО", "Конт ПО")):
            missing_po = combinations["Конт МО"] - combinations["Конт ПО"]
            missing_mo = combinations["Конт ПО"] - combinations["Конт МО"]
            if missing_po:
                errors.append(f"На листе «Конт ПО» отсутствуют комбинации — {len(missing_po)}")
            if missing_mo:
                errors.append(f"На листе «Конт МО» отсутствуют комбинации — {len(missing_mo)}")
    finally:
        workbook.close()
    return {"valid": not errors, "errors": errors[:50], "warnings": warnings, "counts": counts}


@app.get("/api/catalog-management/{business_unit_id}/download")
def catalog_download(business_unit_id: str, request: Request) -> FileResponse:
    _require_catalog_editor(request)
    unit = next((item for item in BUSINESS_UNITS if item["id"] == business_unit_id), None)
    if not unit:
        raise HTTPException(404, "Бизнес-единица не найдена")
    return FileResponse(Path(unit["template"]), filename=Path(unit["template"]).name)


@app.get("/api/catalog-management/{business_unit_id}/versions")
def catalog_versions(business_unit_id: str, request: Request) -> list[dict]:
    _require_catalog_editor(request)
    return list_catalog_versions(business_unit_id)


@app.post("/api/catalog-management/{business_unit_id}/upload")
async def catalog_upload(business_unit_id: str, request: Request, file: UploadFile = File(...)) -> dict:
    editor = _require_catalog_editor(request)
    if not any(item["id"] == business_unit_id for item in BUSINESS_UNITS):
        raise HTTPException(404, "Бизнес-единица не найдена")
    if not file.filename or Path(file.filename).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(400, "Загрузите файл Excel в формате XLSX или XLSM")
    contents = await file.read()
    if len(contents) > 25 * 1024 * 1024:
        raise HTTPException(400, "Excel-файл превышает 25 МБ")
    CATALOG_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    version_id = uuid4().hex
    stored = CATALOG_STORAGE_DIR / f"{version_id}{Path(file.filename).suffix.lower()}"
    stored.write_bytes(contents)
    validation = _validate_catalog_file(stored)
    save_catalog_version({
        "id": version_id,
        "business_unit_id": business_unit_id,
        "filename": file.filename,
        "stored_path": str(stored),
        "created_by": editor["id"],
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "validation": validation,
    })
    return {"id": version_id, "filename": file.filename, "validation": validation}


@app.post("/api/catalog-management/{business_unit_id}/activate/{version_id}")
def catalog_activate(business_unit_id: str, version_id: str, request: Request) -> dict:
    _require_catalog_editor(request)
    version = get_catalog_version(version_id)
    if not version or version["business_unit_id"] != business_unit_id:
        raise HTTPException(404, "Версия справочника не найдена")
    if not version["validation"].get("valid"):
        raise HTTPException(400, "Нельзя применить справочник с ошибками")
    unit = next(item for item in BUSINESS_UNITS if item["id"] == business_unit_id)
    target = Path(unit["template"])
    backup_dir = CATALOG_STORAGE_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    if target.is_file():
        shutil.copy2(target, backup_dir / f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{target.name}")
    temporary = target.with_suffix(target.suffix + ".new")
    shutil.copy2(Path(version["stored_path"]), temporary)
    os.replace(temporary, target)
    APPROVED_CATALOG_HASH_PATH.write_text(file_sha256(target), encoding="utf-8")
    activate_catalog_version(version_id, business_unit_id)
    clear_catalog_cache()
    return {"ok": True}


@app.put("/api/profile/business-unit")
def profile_business_unit_save(preference: BusinessUnitPreference, request: Request) -> dict:
    profile = _current_profile(request)
    profile_id = profile["id"]
    if not profile:
        raise HTTPException(409, "Сначала заполните профиль")
    if not any(item["id"] == preference.business_unit_id for item in BUSINESS_UNITS):
        raise HTTPException(404, "Бизнес-единица не найдена")
    updated = save_last_business_unit(preference.business_unit_id, profile_id)
    return {"last_business_unit_id": updated["last_business_unit_id"] if updated else ""}


@app.post("/api/parse-application")
async def parse_application() -> None:
    # Old browser tabs must refresh; never parse or persist an uploaded questionnaire.
    raise HTTPException(410, "Обновите страницу: PDF-анкеты теперь обрабатываются в браузере")


@app.get("/api/business-units")
def business_units() -> list[dict[str, str]]:
    return [{"id": item["id"], "name": item["name"], "short_code": item["short_code"]} for item in BUSINESS_UNITS]


@app.get("/api/catalog/{business_unit_id}/units")
def catalog_units(business_unit_id: str) -> list[str]:
    try:
        return get_catalog(business_unit_id).list_units()
    except KeyError as exc:
        raise HTTPException(404, "Бизнес-единица не найдена") from exc


@app.get("/api/catalog/{business_unit_id}/jobs")
def catalog_jobs(business_unit_id: str, structural_unit: str) -> list[str]:
    try:
        return get_catalog(business_unit_id).list_jobs(structural_unit)
    except KeyError as exc:
        raise HTTPException(404, "Бизнес-единица не найдена") from exc


@app.get("/api/catalog/{business_unit_id}/requirements")
def requirements(business_unit_id: str, structural_unit: str, job: str) -> dict:
    try:
        catalog = get_catalog(business_unit_id)
    except KeyError as exc:
        raise HTTPException(404, "Бизнес-единица не найдена") from exc
    entry = catalog.get(structural_unit, job)
    if not entry:
        raise HTTPException(404, "Для выбранной комбинации данные не найдены")
    return {
        "factors_text": entry.factors_text,
        "po_required": entry.po_required,
        "po_activities": entry.po_activities,
        "all_po_activities": catalog.all_po_activities,
        "duplicate_resolved": entry.duplicate_resolved,
    }


@app.post("/api/documents")
def create_documents(payload: DocumentRequest, request: Request) -> dict:
    profile = _current_profile(request)
    profile_id = profile["id"]
    if not profile:
        raise HTTPException(409, "Сначала заполните профиль")
    if payload.po_required and not payload.po_activities:
        raise HTTPException(400, "Выберите хотя бы один вид деятельности для ПО")
    try:
        generated = generate_documents(payload, profile)
    except (ValueError, RuntimeError, StopIteration) as exc:
        raise HTTPException(500, str(exc)) from exc

    business_name = next(item["name"] for item in BUSINESS_UNITS if item["id"] == payload.business_unit_id)
    record = {
        "id": generated["id"],
        "created_at": generated["created_at"],
        "candidate_name": payload.full_name,
        "business_unit": business_name,
        "structural_unit": payload.structural_unit,
        "job": payload.job,
        "document_date": payload.document_date,
        "po_required": payload.po_required,
        "formats": payload.formats,
        "files": generated["files"],
        "download_name": generated["download_name"],
        "request": payload.model_dump(),
        "owner_profile_id": profile_id,
    }
    save_document(record)
    return {
        "id": record["id"],
        "download_url": f"/api/archive/{record['id']}/download",
        "download_name": record["download_name"],
    }


@app.get("/api/archive")
def archive_list(request: Request, q: str = "") -> list[dict]:
    return list_documents(_current_profile(request)["id"], q)


@app.get("/api/archive/{document_id}")
def archive_detail(document_id: str, request: Request) -> dict:
    document = get_document(document_id, _current_profile(request)["id"])
    if not document:
        raise HTTPException(404, "Документ не найден")
    return document


@app.delete("/api/archive/{document_id}")
def archive_delete(document_id: str, request: Request) -> dict:
    if not delete_document(document_id, _current_profile(request)["id"]):
        raise HTTPException(404, "Документ не найден")
    return {"ok": True}


@app.get("/api/archive/{document_id}/download")
def archive_download(document_id: str, request: Request) -> FileResponse:
    profile_id = _current_profile(request)["id"]
    document = get_document(document_id, profile_id)
    if not document:
        raise HTTPException(404, "Документ не найден")
    path = resolve_archived_file(document_id, document["download_name"], profile_id)
    if not path:
        raise HTTPException(404, "Файл документа не найден")
    return FileResponse(path, filename=document["download_name"])


@app.get("/api/archive/{document_id}/files/{filename}")
def archive_file(document_id: str, filename: str, request: Request) -> FileResponse:
    path = resolve_archived_file(document_id, filename, _current_profile(request)["id"])
    if not path:
        raise HTTPException(404, "Файл не найден")
    return FileResponse(path, filename=path.name)
