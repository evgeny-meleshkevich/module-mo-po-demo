from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image, ImageChops

from .config import ARCHIVE_DIR, BUSINESS_UNITS
from .integrity import ensure_separate_output, verify_output_template, verify_reference_template
from .models import DocumentRequest


MO_SHEET = "Направления_на_МО"
PO_SHEET = "Направления на ПО (автомат!!!)"
RUSSIAN_LONG_DATE_FORMAT = '[$-0419]d mmmm yyyy "г."'


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._ -]+", "", value).strip().replace(" ", "_")
    return cleaned[:100] or "Направления"


def parse_date(value: str) -> datetime:
    return datetime.strptime(value, "%d.%m.%Y")


def find_soffice() -> str | None:
    configured = os.environ.get("SOFFICE_BIN")
    candidates = [
        configured,
        shutil.which("soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    return next((str(candidate) for candidate in candidates if candidate and Path(candidate).exists()), None)


def detach_template_signature(sheet):
    """Remove the sample signature and retain its exact drawing position."""
    signature = next(
        (image for image in sheet._images if image.width > 300 and image.height < 250),
        None,
    )
    if signature is None:
        return None
    anchor = deepcopy(signature.anchor)
    sheet._images = [image for image in sheet._images if image is not signature]
    return anchor


def add_signature(sheet, signature_path: str, anchor) -> None:
    image = SpreadsheetImage(signature_path)
    if anchor is not None:
        # A two-cell anchor controls both the position and displayed size. Keeping
        # it unchanged preserves the intended slight overlap with the stamp.
        image.anchor = deepcopy(anchor)
    sheet.add_image(image)


def prepare_signature_image(source_path: str, target_path: Path) -> None:
    """Remove only a white background; do not move or resize signature content."""
    source = Image.open(source_path).convert("RGBA")
    rgb = Image.new("RGB", source.size, "white")
    rgb.paste(source.convert("RGB"), mask=source.getchannel("A"))
    difference = ImageChops.difference(rgb, Image.new("RGB", source.size, "white")).convert("L")
    ink_alpha = difference.point(lambda value: 0 if value < 12 else min(255, (value - 12) * 3))
    ink_alpha = ImageChops.multiply(ink_alpha, source.getchannel("A"))
    source.putalpha(ink_alpha)
    bounds = ink_alpha.getbbox()
    if bounds is None:
        raise ValueError("Изображение подписи пустое")
    source.save(target_path, "PNG")


def remove_hidden_mailto(sheet, cell_coordinate: str) -> None:
    """Remove the obsolete hyperlink text that becomes visible after XLSX export."""
    cell = sheet[cell_coordinate]
    cell.value = None
    cell.hyperlink = None


def fill_workbook(
    request: DocumentRequest,
    profile: dict[str, Any],
    output_path: Path,
    document_kind: str,
) -> None:
    if document_kind not in {"mo", "po"}:
        raise ValueError("Неизвестный вид направления")
    if document_kind == "po" and not request.po_required:
        raise ValueError("Нельзя создать направление на ПО со статусом «не требуется»")
    business_unit = next(item for item in BUSINESS_UNITS if item["id"] == request.business_unit_id)
    catalog_path = Path(business_unit["template"])
    verify_reference_template(catalog_path)
    reference_path = Path(business_unit["output_templates"][document_kind])
    verify_output_template(reference_path, document_kind)
    ensure_separate_output(reference_path, output_path)
    workbook = load_workbook(reference_path, keep_vba=False, data_only=True)
    # Materialize cached template text, then keep only the requested form.
    selected_sheet = MO_SHEET if document_kind == "mo" else PO_SHEET
    if selected_sheet not in workbook.sheetnames:
        raise RuntimeError(f"В демонстрационной книге нет листа {selected_sheet}")
    for sheet in list(workbook.worksheets):
        if sheet.title != selected_sheet:
            workbook.remove(sheet)
    workbook.defined_names.clear()
    workbook[selected_sheet].data_validations.dataValidation.clear()
    mo = workbook[MO_SHEET] if MO_SHEET in workbook.sheetnames else None
    po = workbook[PO_SHEET] if PO_SHEET in workbook.sheetnames else None
    date_value = parse_date(request.document_date)
    factors_text = request.factors

    if mo is not None:
        mo["A29"] = request.full_name
        mo["H30"] = parse_date(request.birth_date)
        mo["H30"].number_format = RUSSIAN_LONG_DATE_FORMAT
        mo["H31"] = request.gender
        mo["C33"] = request.oms_number
        mo["A36"] = request.structural_unit
        mo["G37"] = request.job
        mo["A40"] = factors_text
        mo["A42"] = date_value
        mo["A42"].number_format = RUSSIAN_LONG_DATE_FORMAT
        mo["G45"] = "(подпись)"
        mo["I45"] = "(Ф.И.О)"
        remove_hidden_mailto(mo, "F14")
        mo.print_area = "A1:K54"

    if po is not None:
        po["A28"] = request.full_name
        po["H29"] = parse_date(request.birth_date)
        po["H29"].number_format = RUSSIAN_LONG_DATE_FORMAT
        po["H30"] = request.gender
        po["F31"] = "по месту регистрации"
        po["A34"] = request.structural_unit
        po["G35"] = request.job
        po["A37"] = "\n".join(request.po_activities)
        po["A49"] = date_value
        po["A49"].number_format = RUSSIAN_LONG_DATE_FORMAT
        po["G44"] = "(подпись)"
        po["I44"] = "(фамилия, инициалы)"
        remove_hidden_mailto(po, "F15")
        po.print_area = "A1:K51"

    sheets = [sheet for sheet in (mo, po) if sheet is not None]
    signature_anchors = {sheet.title: detach_template_signature(sheet) for sheet in sheets}
    temporary_image_paths: list[Path] = []
    for sheet in sheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.page_setup.scale = None
        sheet.sheet_view.showGridLines = False

    prepared_signature_path: Path | None = None
    if profile.get("role") == "user":
        if mo is not None:
            mo["A44"] = "Специалист по подбору персонала"
            mo["I44"] = profile.get("signer_name", "")
        if po is not None:
            po["A43"] = "Специалист по подбору персонала"
            po["I43"] = profile.get("signer_name", "")
        signature_path = profile.get("signature_path")
        if signature_path and Path(signature_path).is_file():
            with tempfile.NamedTemporaryFile(prefix="medflow-signature-", suffix=".png", delete=False) as temporary:
                prepared_signature_path = Path(temporary.name)
            temporary_image_paths.append(prepared_signature_path)
            prepare_signature_image(signature_path, prepared_signature_path)
            for sheet in sheets:
                add_signature(sheet, str(prepared_signature_path), signature_anchors[sheet.title])
    else:
        if mo is not None:
            mo["A44"] = ""
            mo["I44"] = ""
        if po is not None:
            po["A43"] = ""
            po["I43"] = ""

    workbook.active = 0
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    finally:
        for temporary_path in temporary_image_paths:
            temporary_path.unlink(missing_ok=True)
    verify_reference_template(catalog_path)
    verify_output_template(reference_path, document_kind)


def convert_to_pdf(xlsx_path: Path, output_dir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("Для создания PDF необходимо установить LibreOffice или указать SOFFICE_BIN.")
    with tempfile.TemporaryDirectory(prefix="medflow-lo-") as profile_dir:
        command = [
            soffice,
            f"-env:UserInstallation=file://{profile_dir}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(xlsx_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, timeout=90)
    pdf_path = output_dir / f"{xlsx_path.stem}.pdf"
    if result.returncode != 0 or not pdf_path.is_file():
        details = (result.stderr or result.stdout).strip()
        raise RuntimeError(f"Не удалось создать PDF. {details}".strip())
    return pdf_path


def generate_documents(request: DocumentRequest, profile: dict[str, Any]) -> dict[str, Any]:
    document_id = uuid.uuid4().hex[:12]
    output_dir = ARCHIVE_DIR / document_id
    output_dir.mkdir(parents=True, exist_ok=False)
    business_unit = next(item for item in BUSINESS_UNITS if item["id"] == request.business_unit_id)
    short_code = business_unit["short_code"]
    suffix = safe_filename(request.bundle_name)
    document_kinds = [("mo", "МО")]
    if request.po_required:
        document_kinds.append(("po", "ПО"))
    archived_files: list[Path] = []
    download_files: list[Path] = []
    for kind, label in document_kinds:
        base_name = safe_filename(f"{short_code} Направление на {label} {suffix}")
        workbook_path = output_dir / f"{base_name}.xlsx"
        fill_workbook(request, profile, workbook_path, document_kind=kind)
        pdf_path = convert_to_pdf(workbook_path, output_dir)
        archived_files.extend((workbook_path, pdf_path))
        if "xlsx" in request.formats:
            download_files.append(workbook_path)
        if "pdf" in request.formats:
            download_files.append(pdf_path)

    if len(download_files) == 1:
        download_path = download_files[0]
    else:
        archive_name = safe_filename(f"{short_code} Направления {suffix}")
        download_path = output_dir / f"{archive_name}.zip"
        with zipfile.ZipFile(download_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in download_files:
                archive.write(path, arcname=path.name)

    return {
        "id": document_id,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "files": [path.name for path in archived_files] + ([download_path.name] if len(download_files) > 1 else []),
        "download_name": download_path.name,
        "download_path": download_path,
    }
