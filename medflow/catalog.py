from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from .config import BUSINESS_UNITS
from .integrity import verify_reference_template


SPACE_RE = re.compile(r"\s+")
POINT_RE = re.compile(r"(?:п\.|пункт)\s*\d", re.IGNORECASE)


def normalize(value: object) -> str:
    return SPACE_RE.sub(" ", str(value or "")).strip()


def normalize_key(value: object) -> str:
    return normalize(value).casefold().replace("ё", "е")


def requirement_score(text: str, *, po: bool = False) -> tuple[int, int, int]:
    lowered = text.casefold()
    is_required = 0 if po and "не требуется" in lowered else 1
    points = len(POINT_RE.findall(text))
    return is_required, points, len(text)


def split_po_activities(value: str) -> list[str]:
    if not value or "не требуется" in value.casefold():
        return []
    starts = [match.start() for match in re.finditer(r"(?=Деятельност[ьи])", value)]
    if len(starts) <= 1:
        return [normalize(value)]
    starts.append(len(value))
    return [normalize(value[starts[index] : starts[index + 1]]).rstrip(" ;") for index in range(len(starts) - 1)]


@dataclass(frozen=True)
class CatalogEntry:
    structural_unit: str
    job: str
    factors_text: str
    po_text: str
    duplicate_resolved: bool

    @property
    def po_required(self) -> bool:
        return bool(self.po_text) and "не требуется" not in self.po_text.casefold()

    @property
    def po_activities(self) -> list[str]:
        return split_po_activities(self.po_text)


class Catalog:
    def __init__(self, workbook_path: Path):
        self.workbook_path = workbook_path
        self.entries: dict[tuple[str, str], CatalogEntry] = {}
        self.units: dict[str, str] = {}
        self.jobs_by_unit: dict[str, dict[str, str]] = defaultdict(dict)
        self.all_po_activities: list[str] = []
        self._load()

    def _read_sheet(self, sheet_name: str) -> dict[tuple[str, str], list[tuple[str, str, str]]]:
        workbook = load_workbook(self.workbook_path, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        grouped: dict[tuple[str, str], list[tuple[str, str, str]]] = defaultdict(list)
        for row in range(4, sheet.max_row + 1):
            unit = normalize(sheet.cell(row, 2).value)
            job = normalize(sheet.cell(row, 3).value)
            requirement = str(sheet.cell(row, 4).value or "").strip()
            if unit and job:
                grouped[(normalize_key(unit), normalize_key(job))].append((unit, job, requirement))
        workbook.close()
        return grouped

    def _load(self) -> None:
        verify_reference_template(self.workbook_path)
        mo = self._read_sheet("Конт МО")
        po = self._read_sheet("Конт ПО")
        all_po_values: set[str] = set()

        for key in sorted(set(mo) | set(po)):
            mo_rows = mo.get(key, [])
            po_rows = po.get(key, [])
            if not mo_rows:
                continue
            selected_mo = max(mo_rows, key=lambda row: requirement_score(row[2]))
            selected_po = max(po_rows, key=lambda row: requirement_score(row[2], po=True)) if po_rows else (selected_mo[0], selected_mo[1], "Психиатрическое освидетельствование не требуется")
            unit, job, factors = selected_mo
            po_text = selected_po[2]
            duplicate = len(mo_rows) > 1 or len(po_rows) > 1
            entry = CatalogEntry(unit, job, factors, po_text, duplicate)
            self.entries[key] = entry
            unit_key = normalize_key(unit)
            self.units.setdefault(unit_key, unit)
            self.jobs_by_unit[unit_key].setdefault(normalize_key(job), job)
            all_po_values.update(entry.po_activities)

        self.all_po_activities = sorted(all_po_values, key=str.casefold)

    def list_units(self) -> list[str]:
        return sorted(self.units.values(), key=str.casefold)

    def list_jobs(self, structural_unit: str) -> list[str]:
        jobs = self.jobs_by_unit.get(normalize_key(structural_unit), {})
        return sorted(jobs.values(), key=str.casefold)

    def get(self, structural_unit: str, job: str) -> CatalogEntry | None:
        return self.entries.get((normalize_key(structural_unit), normalize_key(job)))


@lru_cache(maxsize=4)
def get_catalog(business_unit_id: str) -> Catalog:
    unit = next((item for item in BUSINESS_UNITS if item["id"] == business_unit_id), None)
    if not unit:
        raise KeyError(business_unit_id)
    return Catalog(Path(unit["template"]))


def clear_catalog_cache() -> None:
    get_catalog.cache_clear()
