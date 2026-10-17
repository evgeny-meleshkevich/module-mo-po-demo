from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pdfplumber

from medflow.catalog import get_catalog
from medflow.config import BUSINESS_UNITS, ROOT_DIR
from medflow.main import app, _validate_catalog_file, SESSION_COOKIE
import medflow.storage as storage
import medflow.documents as documents
import medflow.main as main


def test_demo_catalog_and_isolation():
    catalog = get_catalog('demo-rt')
    assert len(catalog.entries) == 20
    assert len(catalog.list_units()) == 5
    required = [entry for entry in catalog.entries.values() if entry.po_required]
    assert len(required) == 10
    assert all(1 <= len(entry.po_activities) <= 2 for entry in required)
    assert _validate_catalog_file(BUSINESS_UNITS[0]['template'])['valid']
    assert SESSION_COOKIE == 'medflow_test_session'
    assert ROOT_DIR == Path(__file__).resolve().parents[1]
    assert BUSINESS_UNITS[0]['output_templates']['mo'] == BUSINESS_UNITS[0]['template']
    assert BUSINESS_UNITS[0]['output_templates']['po'] == BUSINESS_UNITS[0]['template']


def test_create_download_and_browser_pdf_boundary(tmp_path, monkeypatch):
    for name, value in {
        'DATABASE_PATH': tmp_path / 'app.db', 'STORAGE_DIR': tmp_path,
        'ARCHIVE_DIR': tmp_path / 'archive', 'SIGNATURE_DIR': tmp_path / 'signatures',
    }.items():
        monkeypatch.setattr(storage, name, value)
    monkeypatch.setattr(documents, 'ARCHIVE_DIR', tmp_path / 'archive')
    monkeypatch.setattr(main, 'SIGNATURE_DIR', tmp_path / 'signatures')
    entry = next(e for e in get_catalog('demo-rt').entries.values() if len(e.po_activities) == 2)
    payload = dict(full_name='Примерный Пример Примерович', birth_date='01.01.1990', gender='Мужской',
                   oms_number='', business_unit_id='demo-rt', structural_unit=entry.structural_unit,
                   job=entry.job, factors=entry.factors_text, po_required=True,
                   po_activities=entry.po_activities, document_date='16.10.2026',
                   formats=['xlsx', 'pdf'], bundle_name='Проверка')
    with TestClient(app) as client:
        assert client.get('/api/business-units').status_code == 401
        assert client.post('/api/auth/setup-admin', json={'email': 'admin@example.invalid', 'password': 'demo-password-123'}).status_code == 200
        assert SESSION_COOKIE in client.cookies
        client.headers['X-CSRF-Token'] = client.get('/api/auth/status').json()['csrf_token']
        assert client.post('/api/parse-application').status_code == 410
        result = client.post('/api/documents', json=payload)
        assert result.status_code == 200, result.text
        assert client.get(result.json()['download_url']).status_code == 200
        folder = tmp_path / 'archive' / result.json()['id']
        workbooks = list(folder.glob('*.xlsx'))
        assert len(workbooks) == 2
        for filename in workbooks:
            workbook = load_workbook(filename)
            assert len(workbook.sheetnames) == 1
            assert not workbook._external_links
            assert not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row)
            assert 'Ресурс Тест' in workbook.active['H9'].value
            workbook.close()
        for filename in folder.glob('*.pdf'):
            with pdfplumber.open(filename) as pdf:
                assert len(pdf.pages) == 1
                assert 'Примерный' in pdf.pages[0].extract_text()
        # A second run without PO must produce only MO.
        result = client.post('/api/documents', json={**payload, 'po_required': False, 'po_activities': []})
        assert result.status_code == 200, result.text
        assert len(list((tmp_path / 'archive' / result.json()['id']).glob('*.xlsx'))) == 1
